"""Scanned PDFs to Word documents that can be edited: reading the layout of the pages."""

import io

import numpy as np
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from rapidfuzz.distance import Indel
from typer.testing import CliRunner

from docfill.api import create_app
from docfill.cli import app as cli
from docfill.config import Settings
from docfill.convert import _for_scans, convert, convert_and_score
from docfill.convert.engines import libreoffice_missing
from docfill.convert.fidelity import words
from docfill.convert.ocr import (
    Line,
    Mark,
    Page,
    Word,
    _apart,
    _colour,
    _even_out,
    _missed,
    _rows,
    _skew,
    _underline,
    layout,
    marks,
    ocr_missing,
    prepare,
    write_docx,
)
from docfill.readers.pdf import scanned_pages
from tests.conftest import make_scanned_pdf, make_text_pdf

requires_tesseract = pytest.mark.skipif(
    ocr_missing(Settings()) is not None, reason="Tesseract OCR is not installed"
)
requires_scan_tools = pytest.mark.skipif(
    ocr_missing(Settings()) is not None or libreoffice_missing() is not None,
    reason="Tesseract OCR and LibreOffice Writer are needed",
)

LEFT, RIGHT, PITCH = 300, 2180, 59  # an A4 page at 300 dpi: margins of 2.5 cm, 12 pt text


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        database_url=f"sqlite:///{tmp_path / 'test.db'}", spacy_model="", output_dir=tmp_path
    )


def line(
    text: str,
    top: int,
    left: int = LEFT,
    right: int | None = None,
    bold: tuple[str, ...] = (),
    height: int = 40,
) -> Line:
    """A line as Tesseract reads it: its words side by side, spread to ``right`` (justified)."""
    tokens = text.split()
    widths = [18 * len(token) for token in tokens]
    gap = 18
    if right is not None and len(tokens) > 1:
        gap = (right - left - sum(widths)) / (len(tokens) - 1)
    found, x = [], float(left)
    for token, width in zip(tokens, widths, strict=True):
        word = Word(token, round(x), top, round(x) + width, top + height, 95.0)
        word.stroke = 2.6 if token in bold or "*" in bold else 1.4
        found.append(word)
        x += width + gap
    return Line(found)


def page(lines: list[Line], number: int = 1) -> Page:
    return Page(number, 2480, 3508, lines)


PARAGRAPH = "Asociatii hotarasc constituirea societatii si semneaza actul de fata"


def first_page() -> Page:
    top = 400
    lines = [
        line("ACT CONSTITUTIV", top, 1040, bold=("*",)),
        line(PARAGRAPH, top + 3 * PITCH, LEFT + 150, RIGHT),  # a first-line indent
        line(PARAGRAPH, top + 4 * PITCH, LEFT, RIGHT),
        line("ultimul rand al paragrafului.", top + 5 * PITCH, LEFT),  # short: it ends there
        line(
            "Art. 1. - Societatea EXEMPLU S.R.L. are sediul in",
            top + 7 * PITCH,
            right=RIGHT,
            bold=("Art.", "1."),
        ),
        line("Municipiul Cluj-Napoca, judetul Cluj.", top + 8 * PITCH),
        line("- numeste si revoca administratorii societatii si stabileste", top + 10 * PITCH),
        line("remuneratia acestora;", top + 11 * PITCH, LEFT + 60),  # under the text, not the dash
        line("- aproba situatiile financiare anuale;", top + 12 * PITCH),
        line(PARAGRAPH, top + 14 * PITCH, LEFT + 150, RIGHT),
        line(PARAGRAPH, top + 15 * PITCH, LEFT, RIGHT),  # goes on on the next page
        line("1", 3300, 1230),  # the page number
    ]
    return page(lines)


def test_paragraphs_are_rebuilt_from_the_place_of_the_lines():
    second = page(
        [
            line("si continua aici, pe pagina urmatoare.", 330),  # the paragraph goes on
            line("Data: 05.10.2026", 330 + 2 * PITCH, LEFT + 150),
            line("2", 3300, 1230),
        ],
        2,
    )
    paragraphs, sheet = layout([first_page(), second])
    texts = [" ".join(line.text for line in p.lines) for p in paragraphs]
    assert texts[0] == "ACT CONSTITUTIV"
    title, first, article, *_ = paragraphs
    assert title.alignment == "center" and all(w.bold for w in title.lines[0].words)
    assert first.alignment == "justify" and len(first.lines) == 3
    assert abs(first.first_indent - 150) < 10 and first.left_indent == 0
    assert first.space_before > PITCH  # the blank line before it
    assert [w.bold for w in article.lines[0].words[:3]] == [True, True, False]
    items = [p for p in paragraphs if p.bullet]
    assert len(items) == 2 and len(items[0].lines) == 2  # the next line of an item stays in it
    assert abs(items[0].left_indent - 36) < 10 and items[0].first_indent < 0  # hanging
    continued = paragraphs[-2]
    assert continued.lines[-1].text.startswith("si continua aici")  # one paragraph, two pages
    assert paragraphs[-1].page_break is False and not any(p.page_break for p in paragraphs)
    assert sheet["size"] == 12  # 59 px at 300 dpi = 14.2 pt, Word's single spacing of 12 pt
    assert sheet["page_numbers"] is True  # and they are not text
    assert not any(text in ("1", "2") for text in texts)
    assert abs(sheet["left"] - LEFT) < 15 and abs(sheet["width"] - sheet["right"] - RIGHT) < 25


def test_a_new_page_of_the_scan_is_a_new_page_of_the_document():
    second = page([line("CAPITOLUL II", 330, 1060, bold=("*",)), line(PARAGRAPH, 420)], 2)
    paragraphs, _ = layout([first_page(), second])
    chapter = next(p for p in paragraphs if p.lines[0].text == "CAPITOLUL II")
    assert chapter.page_break and chapter.alignment == "center"


def test_columns_on_one_row_are_joined_with_a_tab():
    left = line("IONESCU MARIA", 1000)
    right = line("POPESCU ANA", 1004, 1500)
    rows = _rows([left, right, line("Semnatura", 1002, 2000)])
    assert len(rows) == 1 and len(rows[0].breaks) == 2  # three columns, two tabs
    paragraphs, sheet = layout([page([*first_page().lines[:4], *rows])])
    signatures = paragraphs[-1]
    assert len(signatures.tabs) == 2
    document = Document(io.BytesIO(write_docx(paragraphs, sheet)))
    assert document.paragraphs[-1].text == "IONESCU MARIA\tPOPESCU ANA\tSemnatura"


def test_the_word_document_is_written_like_a_typed_one():
    second = page(
        [line("CAPITOLUL II", 330, 1060, bold=("*",)), line(PARAGRAPH, 420), line("2", 3300, 1230)],
        2,
    )
    paragraphs, sheet = layout([first_page(), second])
    document = Document(io.BytesIO(write_docx(paragraphs, sheet, "ACT CONSTITUTIV")))
    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size == Pt(12)
    assert 'w:val="ro-RO"' in normal.element.xml  # spelling checked in Romanian
    section = document.sections[0]
    assert abs(section.left_margin.cm - 2.54) < 0.2  # 300 px at 300 dpi
    written = document.paragraphs
    assert written[0].alignment == WD_ALIGN_PARAGRAPH.CENTER and written[0].runs[0].bold
    assert written[0].paragraph_format.keep_with_next
    assert written[1].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(written[1].paragraph_format.first_line_indent.pt - 36) < 3
    assert written[1].text.count(PARAGRAPH) == 2  # the lines of a paragraph flow, no breaks
    article = written[2]
    assert article.runs[0].text == "Art. 1." and article.runs[0].bold
    assert article.runs[1].bold is None
    item = next(p for p in written if p.text.startswith("-"))
    assert item.text.startswith("-\tnumeste") and item.paragraph_format.first_line_indent.pt < 0
    chapter = next(p for p in written if p.text == "CAPITOLUL II")
    assert chapter.paragraph_format.page_break_before
    assert "PAGE" in document.sections[0].footer.paragraphs[0]._p.xml  # the page number
    assert document.core_properties.title == "ACT CONSTITUTIV"
    assert not document.inline_shapes  # text, no pictures


def test_a_word_cut_on_a_bent_line_is_one_word():
    text = line("sunt de competenta instante", 360)
    end = text.words[-1]
    # Tesseract read the end of the word apart, a little higher: the line bends up
    rest = Line([Word("lor", end.right + 5, end.top - 9, end.right + 59, end.bottom - 9, 86.0)])
    rows = _rows([text, rest])
    assert len(rows) == 1 and rows[0].text == "sunt de competenta instantelor"
    assert not rows[0].breaks  # one line, no column


def test_print_no_word_was_read_from_is_read_again():
    gray = np.full((800, 2480), 255, np.uint8)
    gray[100:140, 300:1000] = 20  # a line read
    gray[100:140, 1030:1060] = 20  # the end of its last word, cut short: no line of its own
    gray[400:440, 300:1500] = 20  # a heading Tesseract took for a picture
    rows = [{"left": "300", "top": "100", "width": "700", "height": "40"}]
    no_stamps = np.zeros(gray.shape, bool)
    again = np.asarray(_missed(gray, no_stamps, rows))
    assert again[420, 800] == 20  # the heading, alone on the page
    assert again[120, 600] == 255 and again[120, 1045] == 255
    rows.append({"left": "300", "top": "400", "width": "1200", "height": "40"})
    assert _missed(gray, no_stamps, rows) is None  # all of it read


def _stamped(lines: list[Line]) -> Image.Image:
    """A page with ``lines`` printed, a blue stamp over the end of the first two, a signature
    in black ink under them, and the red fringe a camera leaves along printed letters."""
    image = Image.new("RGB", (2480, 1400), "white")
    draw = ImageDraw.Draw(image)
    for found in lines:
        for word in found.words:
            box = (word.left, word.top + 8, word.right, word.bottom - 4)
            fringe = (box[0] - 5, box[1] - 5, box[2] + 5, box[3] + 5)
            draw.rectangle(fringe, outline=(210, 120, 120), width=2)
            draw.rectangle(box, fill=(25, 25, 25))
    draw.ellipse((1750, 200, 2150, 600), outline=(40, 70, 200), width=10)
    draw.ellipse((1800, 250, 2100, 550), outline=(40, 70, 200), width=5)
    draw.text((1880, 380), "AVOCAT", fill=(40, 70, 200))
    curve = [(500 + x, 900 + round(60 * np.sin(x / 40))) for x in range(0, 600, 4)]
    draw.line(curve, fill=(60, 60, 70), width=5)
    return image


def test_stamps_and_signatures_are_found_and_print_is_not():
    lines = [line(PARAGRAPH, 300 + PITCH * row, LEFT, RIGHT) for row in range(4)]
    image = _stamped(lines)
    found = marks(np.asarray(image.convert("L")), np.asarray(image), lines)
    assert len(found) == 2  # the stamp, the signature: no print, no fringe
    stamp, signature = sorted(found, key=lambda mark: mark.top)
    assert abs(stamp.left - 1750) < 15 and abs(stamp.top - 200) < 15
    assert abs(stamp.picture.width - 400) < 30 and abs(signature.left - 500) < 15
    pixels = np.asarray(stamp.picture)
    ink = pixels[pixels[..., 3] > 128][:, :3].astype(int)
    assert (ink[:, 2] - ink[:, 0] > 100).mean() > 0.9  # in its blue, the paper transparent
    assert pixels[0, 0, 3] == 0 and pixels[200, 200, 3] == 0
    pixels = np.asarray(signature.picture)
    assert pixels[..., 3].max() > 200 and (pixels[pixels[..., 3] > 128][:, :3] < 90).all()


def test_marks_float_where_they_are_on_the_page():
    first = first_page()
    first.marks = [Mark(1800, 200, Image.new("RGBA", (300, 300), (40, 70, 200, 255)))]
    blank = page([], 2)  # a page with a signature only
    blank.marks = [Mark(600, 900, Image.new("RGBA", (500, 120), (30, 30, 30, 255)))]
    paragraphs, sheet = layout([first, blank])
    assert paragraphs[-1].page is blank and paragraphs[-1].page_break
    document = Document(io.BytesIO(write_docx(paragraphs, sheet, "", [first, blank])))
    assert not document.inline_shapes  # pictures over the text, not in its lines
    xml = document.element.body.xml
    assert xml.count("<wp:anchor ") == 2 and xml.count("<wp:wrapNone/>") == 2
    assert xml.count('relativeFrom="page"') == 4
    assert f"<wp:posOffset>{round(1800 / 300 * 914400)}</wp:posOffset>" in xml
    assert f"<wp:posOffset>{round(900 / 300 * 914400)}</wp:posOffset>" in xml
    held = [p for p in document.paragraphs if "<wp:anchor " in p._p.xml]
    assert held[0].text == "ACT CONSTITUTIV" and held[1].paragraph_format.page_break_before


def test_a_word_cut_at_the_end_of_a_line_is_joined():
    lines = [line(f"{PARAGRAPH} adminis-", 400, LEFT, RIGHT), line("tratorul semneaza.", 459)]
    paragraphs, sheet = layout([page(lines * 1)])
    document = Document(io.BytesIO(write_docx(paragraphs, sheet)))
    assert "adminis-tratorul semneaza." in document.paragraphs[0].text


def test_noise_apart_from_the_text_is_left_out():
    words = line("ACT CONSTITUTIV", 400, 1040).words + [
        Word("Kil", 1800, 400, 1850, 440, 35.0),
        Word("Yoo.", 1900, 400, 1960, 440, 22.0),
    ]
    assert [w.text for w in _apart(words)] == ["ACT", "CONSTITUTIV"]
    sure = words[:2] + [Word("Semnatura", 1800, 400, 1990, 440, 92.0)]
    assert len(_apart(sure)) == 3  # a column read with confidence stays


def test_an_underline_apart_from_the_letters():
    gray = np.full((41, 200), 255, np.uint8)
    for x in range(10, 190, 12):
        gray[2:31, x : x + 4] = 0  # the strokes of the letters
    gray[37:41, :] = 0  # the rule, below a gap
    word = Word("ASOCIATI:", 0, 0, 200, 41, 95.0)
    assert _underline(gray, word) == 10
    assert word.height == 41  # and the height of the letters without it
    word.underline = 10
    assert word.height == 31
    foot = np.full((41, 200), 255, np.uint8)
    for x in range(10, 190, 12):
        foot[2:31, x : x + 4] = 0
    foot[30, :] = 0  # the flat foot of the letters is no underline
    assert _underline(foot, word) == 0


def test_a_scan_is_evened_out_and_straightened():
    image = Image.new("RGB", (1240, 1754), (205, 205, 200))  # grey paper
    draw = ImageDraw.Draw(image)
    for row in range(20):
        draw.rectangle((150, 150 + 60 * row, 1090, 175 + 60 * row), fill=(20, 20, 20))
    draw.ellipse((900, 100, 1100, 300), outline=(40, 60, 200), width=12)  # a blue stamp
    gray = _even_out(image.convert("L"))
    assert np.asarray(gray)[50, 50] == 255 and np.asarray(gray)[160, 400] < 60
    colour = np.asarray(_colour(image)) > 0
    assert colour[200, 906] and not colour[160, 400]
    tilted = image.rotate(1.5, Image.Resampling.BICUBIC, fillcolor=(205, 205, 200))
    assert abs(_skew(_even_out(tilted.convert("L"))) + 1.5) <= 0.2
    straightened, _, colours = prepare(tilted)
    assert abs(_skew(straightened)) <= 0.2
    assert colours.size == straightened.size and colours.mode == "RGB"  # straightened alike


TEXT = ["Asociatul unic hotaraste constituirea societatii", "cu sediul in Cluj-Napoca"]


def test_scanned_pages_are_found():
    assert scanned_pages(make_text_pdf([TEXT, TEXT]), Settings()) == [False, False]
    assert scanned_pages(make_text_pdf([TEXT, []]), Settings()) == [False, False]  # blank
    assert scanned_pages(make_scanned_pdf(["IDENTITY CARD"]), Settings()) == [True]


@requires_tesseract
def test_auto_reads_a_scan_with_ocr(settings):
    scan = make_scanned_pdf(["IDENTITY CARD"])
    assert _for_scans(scan, "auto", settings) == ("ocr", [])
    engine, notes = _for_scans(scan, "pdf2docx", settings)
    assert engine == "pdf2docx" and "Acest PDF este scanat" in notes[0]
    assert _for_scans(make_text_pdf([TEXT]), "auto", settings) == ("auto", [])


# --------------------------------------------------------------------------- a whole scan

ACT = [
    ("ACT CONSTITUTIV", "title"),
    ("al societatii EXEMPLU S.R.L.", "title"),
    (
        "Subsemnatul POPESCU ION, cetatean roman, domiciliat in Municipiul Cluj-Napoca, "
        "strada Florilor numarul 5, judetul Cluj, in calitate de asociat unic, am hotarat "
        "constituirea unei societati cu raspundere limitata, in conditiile prezentului act "
        "constitutiv si ale legii societatilor.",
        "text",
    ),
    ("CAPITOLUL I - DENUMIREA, FORMA JURIDICA SI SEDIUL", "chapter"),
    (
        "Art. 1. - Denumirea societatii este EXEMPLU S.R.L. In toate documentele, facturile, "
        "anunturile si publicatiile emise de societate, denumirea va fi precedata sau urmata "
        "de cuvintele societate cu raspundere limitata si de capitalul social.",
        "text",
    ),
    ("- stabileste strategia si obiectivele de dezvoltare ale societatii;", "item"),
    ("- aproba situatiile financiare anuale si repartizarea profitului;", "item"),
    (
        "Art. 2. - Sediul societatii este in Municipiul Cluj-Napoca, strada Florilor numarul 5, "
        "judetul Cluj. Sediul poate fi mutat in orice alta localitate din Romania, prin "
        "hotararea asociatului unic.",
        "text",
    ),
]


@pytest.fixture(scope="module")
def scanned_act() -> tuple[bytes, list[str]]:
    """An act typed in Word, printed to PDF and scanned: pictures of its pages only."""
    if ocr_missing(Settings()) is not None or libreoffice_missing() is not None:
        pytest.skip("Tesseract OCR and LibreOffice Writer are needed")
    document = Document()
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    style.paragraph_format.space_after = Pt(0)
    for text, kind in ACT:
        paragraph = document.add_paragraph()
        if kind == "title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(text).bold = True
            paragraph.paragraph_format.space_after = Pt(12)
        elif kind == "chapter":
            paragraph.add_run(text).bold = True
            paragraph.paragraph_format.space_before = Pt(12)
        elif kind == "item":
            paragraph.add_run(text)
            paragraph.paragraph_format.left_indent = Cm(1)
            paragraph.paragraph_format.first_line_indent = Cm(-1)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1.27)
            paragraph.paragraph_format.space_before = Pt(12)
            article, _, rest = text.partition(" - ")
            if text.startswith("Art."):
                paragraph.add_run(article).bold = True
                paragraph.add_run(" - " + rest)
            else:
                paragraph.add_run(text)
    buffer = io.BytesIO()
    document.save(buffer)
    printed = convert(buffer.getvalue(), "act.docx").data
    import pypdfium2 as pdfium

    pages = [p.render(scale=150 / 72).to_pil().convert("L") for p in pdfium.PdfDocument(printed)]
    scan = io.BytesIO()
    pages[0].save(scan, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    return scan.getvalue(), words(" ".join(text for text, _ in ACT))


@requires_scan_tools
def test_a_scan_becomes_a_word_document_to_edit(scanned_act, settings):
    scan, expected = scanned_act
    conversion = convert_and_score(scan, "act.pdf", previews=0, settings=settings)
    assert conversion.engine == "ocr"
    document = Document(io.BytesIO(conversion.data))
    assert not document.inline_shapes  # the text itself, not pictures of the pages
    written = [p for p in document.paragraphs if p.text.strip()]
    read = words(" ".join(p.text for p in written))
    common = (len(expected) + len(read) - Indel.distance(expected, read)) // 2
    assert common / len(expected) >= 0.95
    title = written[0]
    assert title.text == "ACT CONSTITUTIV"
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER and title.runs[0].bold
    text = next(p for p in written if p.text.startswith("Subsemnatul"))
    assert text.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(text.paragraph_format.first_line_indent.cm - 1.27) < 0.25
    article = next(p for p in written if p.text.startswith("Art. 1."))
    assert article.runs[0].bold and not article.runs[-1].bold
    items = [p for p in written if p.text.startswith("-")]
    assert len(items) == 2 and items[0].text.startswith("-\t")
    report = conversion.fidelity
    assert report.check("pages").score == 100
    assert report.check("text").score >= 85  # the confidence of the reading
    assert "Citit prin OCR" in report.check("text").detail
    assert report.check("layout").score >= 75
    assert conversion.ocr["words"] >= len(expected) * 0.9


@requires_scan_tools
def test_the_page_converts_a_scan_with_ocr(scanned_act, settings):
    client = TestClient(create_app(settings))
    engines = client.get("/convert/engines").json()["docx"]
    assert [e["name"] for e in engines] == ["pdf2docx", "ocr", "libreoffice"]
    response = client.post(
        "/convert", files={"file": ("scan.pdf", scanned_act[0])}, data={"previews": "1"}
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["engine"] == "ocr" and result["filename"] == "scan.docx"
    assert result["ocr"]["words"] > 50 and result["ocr"]["confidence"] > 85
    text = next(c for c in result["fidelity"]["checks"] if c["key"] == "text")
    assert text["detail"].startswith("Citit prin OCR")


@pytest.fixture(scope="module")
def stamped_act(scanned_act) -> bytes:
    """The scanned act, stamped and signed: a blue stamp beside the title, a signature in black
    ink under the text."""
    import pypdfium2 as pdfium

    pages = [
        p.render(scale=150 / 72).to_pil().convert("RGB") for p in pdfium.PdfDocument(scanned_act[0])
    ]
    draw = ImageDraw.Draw(pages[0])
    draw.ellipse((930, 40, 1130, 240), outline=(40, 70, 200), width=6)
    draw.ellipse((960, 70, 1100, 210), outline=(40, 70, 200), width=3)
    curve = [(260 + x, 1350 + round(30 * np.sin(x / 20))) for x in range(0, 300, 2)]
    draw.line(curve, fill=(50, 50, 60), width=3)
    scan = io.BytesIO()
    pages[0].save(scan, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    return scan.getvalue()


@requires_scan_tools
def test_stamps_and_signatures_are_kept_as_pictures(stamped_act, settings):
    client = TestClient(create_app(settings))
    files = {"file": ("act.pdf", stamped_act)}
    kept = client.post("/convert", files=files, data={"previews": "0"}).json()
    assert kept["ocr"]["marks"] == 2
    assert any("2 ștampile, semnături sau note de mână păstrate" in w for w in kept["warnings"])
    document = Document(io.BytesIO(client.get(kept["url"]).content))
    xml = document.element.body.xml
    assert xml.count("<wp:anchor ") == 2 and not document.inline_shapes
    assert document.paragraphs[0].text == "ACT CONSTITUTIV"  # the text stays text
    layout_check = next(c for c in kept["fidelity"]["checks"] if c["key"] == "layout")
    assert "ștampilele și semnăturile păstrate ca imagini" in layout_check["detail"]
    clean = client.post("/convert", files=files, data={"previews": "0", "marks": "false"}).json()
    assert clean["ocr"]["marks"] == 0
    assert any("omise (o copie curată)" in w for w in clean["warnings"])
    document = Document(io.BytesIO(client.get(clean["url"]).content))
    assert "<wp:anchor " not in document.element.body.xml


@requires_scan_tools
def test_the_cli_makes_a_clean_copy(stamped_act, tmp_path):
    source = tmp_path / "act.pdf"
    source.write_bytes(stamped_act)
    db = ["--db", f"sqlite:///{tmp_path / 'cli.db'}"]
    result = CliRunner().invoke(cli, [*db, "convert", str(source), "--no-marks", "--no-score"])
    assert result.exit_code == 0, result.output
    document = Document(str(tmp_path / "act.docx"))
    assert "<wp:anchor " not in document.element.body.xml
    assert document.paragraphs[0].text == "ACT CONSTITUTIV"


def test_a_line_of_one_word_across_the_page():
    """A dotted line to fill in is one word from margin to margin; the text goes on under it."""
    dots = Line([Word("." * 60, LEFT, 400, RIGHT, 440, 90.0)])
    paragraphs, _ = layout([page([dots, line(PARAGRAPH, 459), line(PARAGRAPH, 518)])])
    assert paragraphs[0].lines[0] is dots and not paragraphs[0].bullet
