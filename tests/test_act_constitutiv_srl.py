"""The act constitutiv of an SRL with a sole associate: the model filers use, filled from the
proof of the firm name, the proof of the registered office, the identity cards of the sole
associate and of the administrator, and the CAEN activities and share capital typed in the
wizard. Every person and company here is fictitious."""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

from docfill.api import create_app
from docfill.app import build_app
from docfill.computed import address_line, with_computed, with_de
from docfill.extraction.articles import extract_premises
from docfill.extraction.derive import complete_values
from docfill.extraction.fields import CONTROL_OPTIONS
from docfill.knowledge import KnowledgeBase
from docfill.pipeline import MEDIA_TYPES, DocFill
from docfill.readers import read_bytes
from docfill.samples import Person, sworn_statement_docx
from docfill.templates import TemplateRepository, bundled_specs_dir, load_directory
from tests.conftest import make_docx
from tests.test_persons import MURESAN
from tests.test_reference_forms import document, statement_values

POPESCU = Person()
ACTIVITIES = (
    "6201 Activități de realizare a software-ului la comandă\n"
    "6202 Activități de consultanță în tehnologia informației\n"
    "6209 Alte activități de servicii privind tehnologia informației"
)
COMPANY = {
    "company_name": "EXEMPLU VERDE S.R.L.",
    "name_reservation_number": "123456",
    "name_reservation_date": "01.09.2026",
    "company_city": "Mun. Timișoara",
    "company_street": "Ale. Teilor, camera 1",
    "company_street_number": "4",
    "company_building": "12",
    "company_floor": "VII",
    "company_apartment": "31",
    "company_county": "Timiș",
    "caen_activities": ACTIVITIES,
    "share_capital": "1.000",
    "share_count": "100",
    "today": "10.09.2026",
}
POPESCU_CLAUSE = (
    f"POPESCU ION-ANDREI, CNP {POPESCU.cnp}, cu domiciliul în Mun. Cluj-Napoca, Str. Florilor "
    "nr. 5, bl. A2, sc. 1, et. 3, ap. 10, jud. Cluj, țara România, cetățenia Română, născut în "
    "Mun. Sibiu, jud. Sibiu, țara România, la data de 14.11.1987, identificat prin CI, seria AX, "
    "nr. 123456, emisă de SPCLEP Cluj-Napoca, la data de 22.06.2022, valabilă până la data de "
    "14.11.2032"
)


def act_text(docx: bytes) -> str:
    """The text of the act, a Word document, paragraph after paragraph."""
    paragraphs = Document(io.BytesIO(docx)).paragraphs
    return " ".join(" ".join(paragraph.text for paragraph in paragraphs).split())


def fill(values: dict[str, str], settings) -> str:
    result = DocFill(settings).fill(document("act-constitutiv-srl"), None, values)
    assert result.missing == []
    assert result.output == "docx"
    return act_text(result.data)


def test_the_act_copies_the_model_verbatim():
    """Every paragraph of the template without blanks is the model's own text, and every
    article of the model is in the template."""
    model_path = bundled_specs_dir() / "model-act-constitutiv-srl-asociat-unic.docx"
    model = " ".join(read_bytes(model_path.read_bytes(), model_path.name).text.split())
    body = document("act-constitutiv-srl").body or ""
    verbatim = [
        line.lstrip("# ").strip()
        for line in body.splitlines()
        if line.strip() and not line.lstrip().startswith("[[") and "{{" not in line
    ]
    assert len(verbatim) > 100
    assert [line for line in verbatim if " ".join(line.split()) not in model] == []
    assert set(re.findall(r"Art\. \d+(?:\.\d+)?\.", model)) <= set(
        re.findall(r"Art\. \d+(?:\.\d+)?\.", body)
    )


def test_a_sole_associate_who_administers_the_company(settings):
    values = {
        **statement_values(POPESCU),  # read from the identity card
        **COMPANY,
        "associate": "x",
        "board_role": "administrator unic",
    }
    text = fill(values, settings)
    assert "ACT CONSTITUTIV al Societății EXEMPLU VERDE S.R.L. Asociat unic: " in text
    assert f"Asociat unic: {POPESCU_CLAUSE}." in text
    assert (
        "înființarea Societății EXEMPLU VERDE S.R.L.— societate cu răspundere limitată/S.R.L."
        in text
    )
    assert (
        "Art. 1.2. — Denumirea societății este: EXEMPLU VERDE S.R.L.— societate cu răspundere "
        "limitată/S.R.L., conform dovezii privind disponibilitatea firmei nr. 123456 din "
        "01.09.2026, eliberată de Oficiul Național al Registrului Comerțului."
    ) in text
    assert (
        "Art. 1.4. — Sediul societății este în Mun. Timișoara, Ale. Teilor nr. 4, bl. 12, "
        "et. VII, ap. 31, camera 1, jud. Timiș."
    ) in text
    # the object of activity: the main CAEN activity unless written otherwise
    assert (
        "Art. 2.1. — Obiectul de activitate al societății este: Activități de realizare a "
        "software-ului la comandă Domeniul principal de activitate corespunde grupei CAEN 620, "
        "căruia îi corespunde clasa CAEN 6201. — activitatea principală clasa CAEN 6201 și "
        "denumirea activității Activități de realizare a software-ului la comandă — activități "
        "secundare: — clasa CAEN 6202 și denumirea activității Activități de consultanță în "
        "tehnologia informației; — clasa CAEN 6209 și denumirea activității Alte activități de "
        "servicii privind tehnologia informației."
    ) in text
    assert (
        "capitalul social subscris al societății este de 1.000 lei, aport în numerar, fiind "
        "împărțit într-un număr de 100 de părți sociale, cu o valoare nominală de 10 lei/parte "
        "socială."
    ) in text
    assert (
        f"Capitalul social este deținut de către unicul asociat astfel: {POPESCU_CLAUSE}, deține "
        "o participație de 100% din capitalul social al societății."
    ) in text
    assert (
        f"Administrarea societății se face de către: {POPESCU_CLAUSE}, pe perioadă de 30 de "
        "ani, cu posibilitatea reînnoirii mandatului."
    ) in text
    assert f"beneficiarul real al societății este: {POPESCU_CLAUSE}, deține o participație" in text
    assert "intră în vigoare astăzi, 10.09.2026, data semnării" in text
    assert "Data: 10.09.2026 Asociat Unic: POPESCU ION-ANDREI" in text


def test_the_associate_and_the_administrator_are_two_persons(settings):
    values = {
        **statement_values(MURESAN),
        "board_role": "administrator unic",
        **statement_values(POPESCU, "p2_"),
        "p2_associate": "x",
        **COMPANY,
        "activity_object": "Producția de software",
    }
    text = fill(values, settings)
    assert f"Asociat unic: {POPESCU_CLAUSE}." in text
    assert "Administrarea societății se face de către: MUREȘAN VLAD, CNP" in text
    assert "beneficiarul real al societății este: POPESCU ION-ANDREI" in text
    assert "Asociat Unic: POPESCU ION-ANDREI" in text
    assert "Obiectul de activitate al societății este: Producția de software" in text


def test_the_act_is_a_word_document_laid_out_like_the_model(settings):
    values = {**statement_values(POPESCU), **COMPANY, "associate": "x"}
    values["board_role"] = "administrator unic"
    result = DocFill(settings).fill(document("act-constitutiv-srl"), None, values)
    assert (result.output, result.suffix, result.data[:2]) == ("docx", ".docx", b"PK")
    word = Document(io.BytesIO(result.data))
    normal = word.styles["Normal"].font
    assert (normal.name, normal.size.pt) == ("Times New Roman", 12)
    section = word.sections[0]
    assert (section.page_width.inches, section.left_margin.inches) == (8.5, 1)
    paragraphs = {paragraph.text: paragraph for paragraph in word.paragraphs}

    def runs(text: str) -> list[tuple[str, bool, bool]]:
        found = next(p for t, p in paragraphs.items() if t.startswith(text))
        return [(run.text, bool(run.bold), bool(run.italic)) for run in found.runs]

    title = paragraphs["ACT CONSTITUTIV"]
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER and title.runs[0].font.size.pt == 14
    assert runs("al Societății") == [("al Societății EXEMPLU VERDE S.R.L.", True, False)]
    assert paragraphs["CAPITOLUL I"].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert runs("CAPITOLUL I") == [("CAPITOLUL I", True, False)]
    article = runs("Art. 1.3.")
    assert article == [
        ("Art. 1.3.", True, False),
        (" — Durata de funcționare a societății este nedeterminată.", False, False),
    ]
    assert paragraphs[next(t for t in paragraphs if t.startswith("Art. 1.4."))].alignment == (
        WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    # the firm, the associate and the administrator in bold, as in the model
    assert runs(POPESCU_CLAUSE[:20]) == [(POPESCU_CLAUSE, True, False), (".", False, False)]
    assert runs("Asociat Unic:") == [
        ("Asociat Unic: ", False, False),
        ("POPESCU ION-ANDREI", True, False),
    ]
    assert runs("(semnătura)") == [("(semnătura)", False, True)]
    assert word.core_properties.title == "Act constitutiv SRL - asociat unic"


def test_the_other_documents_stay_pdf(settings):
    values = {**statement_values(POPESCU), **COMPANY, "board_role": "administrator unic"}
    result = DocFill(settings).fill(document("declaratie-administrator"), None, values, True)
    assert (result.output, result.suffix, result.data[:4]) == ("pdf", ".pdf", b"%PDF")
    assert result.pdf == result.data


def test_without_secondary_activities_their_heading_is_left_out(settings):
    values = {
        **statement_values(POPESCU),
        **COMPANY,
        "caen_activities": ACTIVITIES.split("\n")[0],
    }
    text = fill(values, settings)
    assert "activități secundare" not in text
    assert "Asociat unic: POPESCU ION-ANDREI" in text  # nobody marked: person 1 is the associate


def test_the_values_the_act_composes():
    assert [with_de(n) for n in (1, 19, 20, 50, 100, 101, 120)] == [
        "1", "19", "20 de", "50 de", "100 de", "101", "120 de",
    ]  # fmt: skip
    office = {
        "company_city": "Mun. Timișoara",
        "company_street": "Ale. Teilor, camera 1",
        "company_street_number": "4",
        "company_apartment": "31",
        "company_county": "Timiș",
    }
    assert address_line(office, "company_", county="county") == (
        "Mun. Timișoara, Ale. Teilor nr. 4, ap. 31, camera 1, jud. Timiș"
    )
    computed = with_computed({"caen_activities": ACTIVITIES, "share_capital": "500"})
    assert (computed["main_caen_group"], computed["main_caen_class"]) == ("620", "6201")
    assert computed["secondary_activity_lines"].splitlines()[-1].endswith("informației.")
    # the object of activity is proposed from the main activity, and can be written otherwise
    proposed = complete_values({"caen_activities": ACTIVITIES})["activity_object"][0]
    assert proposed == "Activități de realizare a software-ului la comandă"


def test_the_roles_the_act_needs_are_derived():
    derived = complete_values(
        {
            "last_name": "Popescu",
            "first_name": "Ion",
            "associate": "x",
            "board_role": "administrator unic",
            "share_count": "50",
            "company_name": "EXEMPLU VERDE S.R.L.",
        }
    )
    assert derived["shares"][0] == "50"  # the sole associate holds every share
    assert derived["capacity"][0] == "asociat unic și administrator"
    assert derived["associates"][0] == "POPESCU ION | 50 părți sociale (100%)"
    assert complete_values({"shares": "10"})["associate"][0] == "x"


# --------------------------------------------------------------------------- from the documents


def test_the_office_from_a_scanned_comodat():
    # what OCR reads on a scan: "lancu" for Iancu, "BI." for bl., no diacritics in the town
    text = (
        "CONTRACT DE COMODAT\nCedarea folosinței pentru imobilul situat în Jud. Timiș, Mun. "
        "Timisoara, Ale. lancu Jianu, Nr. 4, BI. 12, Et. VII, Ap. 31,\nCamera 1. Spațiul va fi "
        "utilizat cu destinația de sediu social.\n"
    )
    found = {field.name: field.value for field in extract_premises(text)}
    derived = complete_values({"company_address": found["company_address"]})
    office = {name: value for name, (value, _) in derived.items()}
    assert address_line(office, "company_", county="county") == (
        "Mun. Timișoara, Ale. Iancu Jianu nr. 4, bl. 12, et. VII, ap. 31, camera 1, jud. Timiș"
    )


def _context(settings):
    context = build_app(settings, use_ner=False)
    context.knowledge.seed()
    with context.sessions() as session:
        repository = TemplateRepository(session)
        for spec in load_directory(bundled_specs_dir()):
            repository.save(spec)
    return context


def test_the_act_is_offered_for_an_srl_only(settings):
    knowledge: KnowledgeBase = _context(settings).knowledge
    forms = {
        key: [form.template for form in knowledge.procedure(key).forms]
        for key in ("srl.infiintare", "srl-d.infiintare", "sa.infiintare")
    }
    assert "act-constitutiv-srl" in forms["srl.infiintare"]
    assert "act-constitutiv-srl" not in forms["srl-d.infiintare"]
    assert "act-constitutiv-srl" not in forms["sa.infiintare"]
    assert "act-constitutiv-sa" in forms["sa.infiintare"]


def test_wizard_takes_the_associate_and_the_administrator_apart(settings, tmp_path):
    _context(settings)
    settings = settings.model_copy(update={"output_dir": tmp_path / "out"})
    popescu = ("popescu.docx", sworn_statement_docx(POPESCU), "application/octet-stream")
    muresan = ("muresan.docx", sworn_statement_docx(MURESAN), "application/octet-stream")
    templates = ["act-constitutiv-srl"]
    with TestClient(create_app(settings)) as client:
        # the sole associate administers the company
        body = client.post(
            "/wizard/analyze", files=[("associate", popescu)], data={"templates": templates}
        ).json()
        rows = {row["name"]: row["value"] for row in body["rows"]}
        assert [d["person"] for d in body["documents"]] == [1]
        assert rows["last_name"] == "Popescu" and rows["associate"] == "x"
        assert rows["board_role"] == "administrator unic"
        # also the beneficial owner (for the declaration; the act writes it from the associate)
        others = {field["name"]: field["value"] for field in body["other_fields"]}
        assert others["beneficial_owner"] == CONTROL_OPTIONS[0]
        # another person administers it: the administrator is person 1 (signs the requests)
        body = client.post(
            "/wizard/analyze",
            files=[("associate", popescu), ("administrator", muresan)],
            data={"templates": templates},
        ).json()
        rows = {row["name"]: row["value"] for row in body["rows"]}
        assert [d["person"] for d in body["documents"]] == [2, 1]
        assert (rows["last_name"], rows["board_role"], rows["associate"]) == (
            "Mureșan",
            "administrator unic",
            "",
        )
        assert (rows["p2_last_name"], rows["p2_associate"]) == ("Popescu", "x")
        others = {field["name"]: field["value"] for field in body["other_fields"]}
        assert others["p2_beneficial_owner"] == CONTROL_OPTIONS[0]
        values = {name: value for name, value in rows.items() if value}
        values.update({**COMPANY, "share_count": "100"})
        exported = client.post(
            "/wizard/export",
            json={"templates": templates, "values": values, "allow_missing": True},
        ).json()
        (saved,) = exported["files"]
        assert saved["filename"].startswith("act-constitutiv-srl_EXEMPLU_VERDE")
        assert saved["filename"].endswith(".docx")  # the act is written in Word
        download = client.get(saved["url"])
        assert download.headers["content-type"] == MEDIA_TYPES["docx"]
        assert download.content == (tmp_path / "out" / saved["filename"]).read_bytes()
        text = act_text(download.content)
        assert f"Asociat unic: {POPESCU_CLAUSE}." in text
        assert "Administrarea societății se face de către: MUREȘAN VLAD" in text


def test_the_act_from_the_proofs_of_the_name_and_the_office(settings):
    context = _context(settings)
    docfill = context.docfill
    proof = make_docx(
        [
            "MINISTERUL JUSTIŢIEI",
            "DOVADĂ",
            "privind disponibilitatea şi rezervarea denumirii firmei",
            "Nr.: 123456/01.09.2026",
            "s-a procedat la verificarea",
            "disponibilităţii denumirii firmei EXEMPLU VERDE S.R.L. solicitată de",
        ]
    )
    comodat = make_docx(
        [
            "CONTRACT DE COMODAT DIN 09.09.2026",
            "Cedarea folosinței pentru imobilul situat în Jud. Timiș, Mun. Timisoara, Ale. lancu "
            "Jianu, Nr. 4, BI. 12, Et. VII, Ap. 31, Camera 1. Spațiul va fi sediu social.",
        ]
    )
    analyses = [
        docfill.analyze_bytes(proof, "dovada.docx"),
        docfill.analyze_bytes(comodat, "comodat.docx"),
        docfill.analyze_bytes(sworn_statement_docx(POPESCU), "ci.docx", "declaratie_administrator"),
    ]
    combined = docfill.combine(analyses)
    values = {name: field.value for name, field in combined.fields.items()}
    values.update(caen_activities=ACTIVITIES, share_capital="1.000", share_count="100")
    text = act_text(docfill.fill(document("act-constitutiv-srl"), None, values, True).data)
    assert "Denumirea societății este: EXEMPLU VERDE S.R.L.— societate" in text
    assert "disponibilitatea firmei nr. 123456 din 01.09.2026" in text
    assert (
        "Sediul societății este în Mun. Timișoara, Ale. Iancu Jianu nr. 4, bl. 12, et. VII, "
        "ap. 31, camera 1, jud. Timiș."
    ) in text
    assert f"Asociat unic: {POPESCU_CLAUSE}." in text
