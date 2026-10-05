"""Converting documents (Word → PDF, PDF → Word) and scoring how faithful the result is."""

import io
import json
import zipfile

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from typer.testing import CliRunner

from docfill.api import create_app
from docfill.cli import app as cli
from docfill.config import Settings
from docfill.convert import (
    compare_files,
    convert,
    convert_and_score,
    document_kind,
    engines_for,
)
from docfill.convert.engines import choose_engine, libreoffice_missing, pdf2docx_missing
from docfill.convert.fidelity import (
    Rendition,
    compare,
    font_key,
    font_name,
    match_font,
    page_similarity,
    pdf_rendition,
    twin,
    word_rendition,
    words,
)
from docfill.errors import (
    ConversionError,
    ConverterUnavailableError,
    DocumentReadError,
    UnsupportedDocumentError,
)
from tests.conftest import make_image, make_scanned_pdf, make_text_pdf

requires_libreoffice = pytest.mark.skipif(
    libreoffice_missing() is not None, reason="LibreOffice Writer is not installed"
)
requires_pdf2docx = pytest.mark.skipif(
    pdf2docx_missing() is not None or libreoffice_missing() is not None,
    reason="pdf2docx (and LibreOffice, to lay the Word document out) are not installed",
)

LETTER = [
    ["Declaratie pe propria raspundere", "Subsemnatul POPESCU Ion, CNP 1800101123456,"],
    ["Pagina a doua", "Data: 05.10.2026"],
]


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        database_url=f"sqlite:///{tmp_path / 'test.db'}",
        spacy_model="",
        output_dir=tmp_path / "out",
    )


def word_document(fonts: dict[str, str] | None = None) -> bytes:
    """A Word document: a heading, a paragraph in Arial, one in the default font."""
    document = Document()
    document.add_heading("Act adițional", 1)
    run = document.add_paragraph().add_run("Societatea EXEMPLU S.R.L., Mun. Timișoara")
    run.font.name = (fonts or {}).get("run", "Arial")
    document.add_paragraph("Asociatul unic ȘTEFĂNESCU Ion hotărăște.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def with_statistics(data: bytes, pages: int, words_saved: int) -> bytes:
    """The document as Word saves it: with the number of pages and words it counted."""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/app.xml":
                content = (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'extended-properties"><Pages>'
                    f"{pages}</Pages><Words>{words_saved}</Words>"
                    "<Application>Microsoft Office Word</Application></Properties>"
                ).encode()
            target.writestr(item, content)
    return buffer.getvalue()


# --------------------------------------------------------------------------- the pieces


def test_words_are_compared_without_ligatures_and_soft_hyphens():
    assert words("ﬁșier inter­național, Bra\u0002 șov") == ["fișier", "internațional", "Brașov"]


@pytest.mark.parametrize(
    ("name", "family", "key"),
    [
        ("ABCDEF+TimesNewRomanPS-BoldMT", "Times New Roman", "timesnewroman"),
        ("TimesNewRomanPSMT", "Times New Roman", "timesnewroman"),
        ("Times New Roman", "Times New Roman", "timesnewroman"),
        ("Arial-BoldMT", "Arial", "arial"),
        ("Calibri,Bold", "Calibri", "calibri"),
        ("BAAAAA+LiberationSerif-Italic", "Liberation Serif", "liberationserif"),
        ("DejaVuSans", "DejaVu Sans", "dejavusans"),
        ("Arial Narrow", "Arial Narrow", "arialnarrow"),
    ],
)
def test_font_names(name, family, key):
    assert font_name(name) == family
    assert font_key(name) == key


def test_fonts_of_the_same_widths():
    available = {font_key(n): n for n in ("Liberation Serif", "Carlito", "DejaVu Sans")}
    assert match_font("DejaVuSans-Bold", available) == ("same", "DejaVu Sans")
    assert match_font("Times New Roman", available) == ("metric", "Liberation Serif")
    assert match_font("Calibri", available) == ("metric", "Carlito")
    assert match_font("Verdana", available) == ("replaced", None)
    assert twin("Calibri") == "Carlito"
    assert twin("Arial Narrow") == "Liberation Sans Narrow"
    assert twin("Verdana") is None


def test_page_similarity_tolerates_a_point_not_a_line():
    page = np.zeros((200, 150), bool)
    page[40:48, 20:130] = True  # a line of text
    page[60:68, 20:100] = True
    assert page_similarity(page, page) == 1.0
    assert page_similarity(page, np.roll(page, 1, axis=1)) == 1.0  # within the tolerance
    assert page_similarity(page, np.roll(page, 20, axis=0)) < 0.5  # one line lower
    blank = np.zeros_like(page)
    assert page_similarity(blank, blank) == 1.0
    assert page_similarity(page, blank) == 0.0
    half = page.copy()
    half[60:68] = False  # the second line is missing
    assert 0.7 < page_similarity(page, half) < 0.9


def test_fonts_of_a_word_document_as_word_picks_them():
    rendition = word_rendition(word_document(), "docx", 60, layout=False)
    assert rendition.words[:2] == ["Act", "adițional"]
    assert not rendition.words_from_pages and rendition.pdf is None
    # the run's own font, then the theme's fonts of the styles (python-docx's template: Cambria
    # for the text, Calibri for the headings)
    assert rendition.fonts == {
        "Calibri": len("Act adițional"),
        "Arial": len("Societatea EXEMPLU S.R.L., Mun. Timișoara"),
        "Cambria": len("Asociatul unic ȘTEFĂNESCU Ion hotărăște."),
    }


def test_the_number_of_pages_word_saved_is_used_only_when_up_to_date():
    data = word_document()
    # python-docx keeps the statistics of its template (Words 0): out of date
    assert word_rendition(data, "docx", 60, layout=False).pages is None
    counted = len(word_rendition(data, "docx", 60, layout=False).words)
    current = word_rendition(with_statistics(data, 3, counted), "docx", 60, layout=False)
    assert current.pages == 3
    assert "Microsoft Office Word" in current.pages_source
    stale = word_rendition(with_statistics(data, 3, 500), "docx", 60, layout=False)
    assert stale.pages is None


def test_identical_pdfs_are_1_to_1():
    pdf = make_text_pdf(LETTER)
    report = compare(pdf_rendition(pdf), pdf_rendition(pdf), previews=1)
    assert report.score == 100
    assert report.verdict[0] == "identical"
    assert [page.score for page in report.pages] == [100, 100]
    assert report.pages[0].diff.startswith("data:image/png;base64,")
    assert report.pages[1].diff is None  # only the first page as a picture
    assert report.check("images").score is None  # no pictures: not measured


def test_a_missing_line_and_an_extra_page_are_found():
    original = pdf_rendition(make_text_pdf(LETTER))
    converted = pdf_rendition(make_text_pdf([LETTER[0][:1], LETTER[1], ["De adaugat"]]))
    report = compare(original, converted, previews=0)
    text, pages, layout = report.check("text"), report.check("pages"), report.check("layout")
    assert pages.score == 50  # 3 pages for 2
    assert [page.score for page in report.pages][1:] == [100, 0]
    assert layout.score < 70 and layout.status == "bad"
    assert text.score < 90
    missing = [d["missing"] for d in text.extra["differences"]]
    assert "Subsemnatul POPESCU Ion CNP 1800101123456" in missing
    assert text.extra["differences"][-1]["added"] == "De adaugat"
    assert report.verdict[0] == "different"
    data = report.as_dict()
    assert data["checks"][0]["title"] == "Layout, page by page"
    assert json.dumps(data)  # plain JSON


def test_without_pages_the_score_says_so():
    original = Rendition("docx", None, ["Un", "text"], False, {"Times New Roman": 8}, 0, 1)
    converted = Rendition("pdf", b"", ["Un", "text", "1"], True, {"LiberationSerif": 1}, 0, 1)
    report = compare(original, converted)
    assert report.check("layout").score is None
    assert report.check("text").score == 100  # the page number printed is not held against it
    assert report.check("fonts").score == 90  # Liberation Serif: the widths of Times New Roman
    assert report.as_dict()["label"].endswith("(pages not compared)")


def test_a_scan_has_no_text_to_compare():
    rendition = pdf_rendition(make_scanned_pdf(["IDENTITY CARD"]))
    assert rendition.words == [] and rendition.images == 1
    assert "no text" in rendition.warnings[0]
    report = compare(rendition, rendition)
    assert report.check("text").score is None
    assert report.check("images").score == 100
    assert report.warnings == rendition.warnings


def test_only_word_and_pdf_documents_are_converted(settings):
    with pytest.raises(UnsupportedDocumentError):
        document_kind(make_image(["x"]), "card.png", settings)
    with pytest.raises(DocumentReadError):
        document_kind(b"", "empty.pdf", settings)
    small = settings.model_copy(update={"max_file_size": 10})
    with pytest.raises(DocumentReadError):
        document_kind(make_text_pdf(LETTER), "big.pdf", small)
    assert document_kind(make_text_pdf(LETTER), "a.pdf", settings) == "pdf"
    assert document_kind(word_document(), "a.docx", settings) == "docx"
    with pytest.raises(ConversionError, match="already a PDF"):
        convert(make_text_pdf(LETTER), "a.pdf", "pdf", settings=settings)
    with pytest.raises(ConversionError, match="choose pdf or docx"):
        convert(make_text_pdf(LETTER), "a.pdf", "odt", settings=settings)


def test_engines_and_what_they_need(monkeypatch, tmp_path):
    assert [e.name for e in engines_for("pdf")] == ["libreoffice"]
    assert [e.name for e in engines_for("docx")] == ["pdf2docx", "libreoffice"]
    with pytest.raises(ConversionError, match="choose auto, libreoffice"):
        choose_engine("pdf", "pdf2docx")
    # LibreOffice without Writer (a distribution's core package alone)
    program = tmp_path / "program"
    program.mkdir()
    (program / "soffice").write_text("")
    (program / "soffice.bin").write_text("")
    monkeypatch.setattr("docfill.convert.engines.libreoffice", lambda: str(program / "soffice"))
    assert "Writer is not installed" in libreoffice_missing()
    monkeypatch.setattr("docfill.convert.engines.libreoffice", lambda: None)
    assert "LibreOffice is not installed" in libreoffice_missing()
    with pytest.raises(ConverterUnavailableError, match="no engine can convert to pdf"):
        choose_engine("pdf")
    assert engines_for("pdf")[0].as_dict()["available"] is False


# --------------------------------------------------------------------------- conversions


@requires_libreoffice
def test_word_to_pdf(settings):
    data = with_statistics(word_document({"run": "Times New Roman"}), 1, 15)
    conversion = convert_and_score(data, "act.docx", settings=settings)
    assert conversion.engine == "libreoffice" and conversion.target == "pdf"
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(conversion.data)).pages)
    assert "ȘTEFĂNESCU" in text and "Timișoara" in text
    report = conversion.fidelity
    assert report.check("text").score == 100
    assert report.check("pages").score == 100  # 1 page, as Word counted
    assert report.check("layout").score is None  # LibreOffice would be compared with itself
    assert report.check("fonts").score is not None
    assert report.score >= 90


@requires_libreoffice
def test_word_to_pdf_against_the_real_pdf(settings):
    data = word_document()
    made = convert(data, "act.docx", settings=settings).data
    conversion = convert_and_score(
        data, "act.docx", reference=("real.pdf", made), settings=settings
    )
    assert conversion.reference.score == 100  # the same fonts, the same pages
    assert conversion.reference.check("layout").score == 100
    with pytest.raises(ConversionError, match="must be a PDF document"):
        convert_and_score(data, "act.docx", reference=("real.docx", data), settings=settings)


@requires_pdf2docx
def test_pdf_to_word(settings):
    conversion = convert_and_score(make_text_pdf(LETTER), "decl.pdf", settings=settings)
    assert conversion.engine == "pdf2docx" and conversion.target == "docx"
    paragraphs = [p.text for p in Document(io.BytesIO(conversion.data)).paragraphs]
    assert any("POPESCU Ion" in text for text in paragraphs)
    report = conversion.fidelity
    assert report.check("text").score == 100
    assert report.check("pages").score == 100
    assert report.check("layout").score >= 90
    assert report.score >= 90


@requires_libreoffice
def test_pdf_to_word_with_libreoffice(settings):
    conversion = convert_and_score(
        make_text_pdf(LETTER), "decl.pdf", engine="libreoffice", previews=0, settings=settings
    )
    assert conversion.engine == "libreoffice"
    assert "frame" in conversion.warnings[0]
    assert conversion.fidelity.check("text").score == 100
    assert conversion.fidelity.check("layout").score is not None


@requires_libreoffice
def test_compare_a_word_document_with_its_pdf(settings):
    data = word_document()
    pdf = convert(data, "act.docx", settings=settings).data
    report = compare_files(("act.docx", data), ("act.pdf", pdf), previews=0, settings=settings)
    assert report.check("layout").score == 100
    assert report.check("text").score == 100
    # Arial, Calibri, Cambria → Liberation Sans, Carlito, Caladea (unless they are installed)
    assert report.check("fonts").score >= 90
    assert report.against == "reference"


# --------------------------------------------------------------------------- page, API, CLI


@pytest.fixture
def client(settings):
    return TestClient(create_app(settings))


def test_convert_page_and_engines(client):
    page = client.get("/convert")
    assert page.status_code == 200 and "Convert" in page.text
    assert set(client.get("/convert/engines").json()) == {"pdf", "docx"}
    assert set(client.get("/health").json()["convert"]) == {"pdf", "docx"}
    assert 'href="/convert"' in client.get("/wizard").text


def test_compare_endpoint(client):
    pdf = make_text_pdf(LETTER)
    other = make_text_pdf([LETTER[0]])
    response = client.post(
        "/convert/compare",
        files={"real": ("real.pdf", pdf), "converted": ("converted.pdf", other)},
        data={"previews": "1"},
    )
    assert response.status_code == 200
    report = response.json()
    assert report["checks"][2]["score"] == 50  # 1 page for 2
    assert report["pages"][0]["diff"] and report["pages"][1]["diff"] is None
    assert report["verdict"] == "different"


def test_convert_refuses_what_it_cannot_convert(client):
    image = client.post("/convert", files={"file": ("card.png", make_image(["x"]))})
    assert image.status_code == 415
    same = client.post(
        "/convert", files={"file": ("a.pdf", make_text_pdf(LETTER))}, data={"to": "pdf"}
    )
    assert same.status_code == 422
    assert client.get("/convert/files/..%2Fsecret.pdf").status_code == 404
    assert client.get("/convert/files/missing.pdf").status_code == 404


@requires_libreoffice
def test_convert_endpoint(client, settings):
    response = client.post(
        "/convert",
        files={"file": ("Act adițional.docx", word_document())},
        data={"previews": "0"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["filename"] == "Act_aditional.pdf"
    assert result["fidelity"]["checks"][1]["score"] == 100
    assert result["reference"] is None
    saved = settings.output_dir / "converted" / "Act_aditional.pdf"
    assert saved.is_file()
    download = client.get(result["url"])
    assert download.status_code == 200 and download.content == saved.read_bytes()
    wrong = client.post(
        "/convert",
        files={"file": ("a.docx", word_document()), "reference": ("b.docx", word_document())},
    )
    assert wrong.status_code == 422


runner = CliRunner()


def test_cli_compare(tmp_path):
    real, same, other = (tmp_path / name for name in ("real.pdf", "same.pdf", "other.pdf"))
    real.write_bytes(make_text_pdf(LETTER))
    same.write_bytes(make_text_pdf(LETTER))
    other.write_bytes(make_text_pdf([LETTER[0]]))
    db = ["--db", f"sqlite:///{tmp_path / 'cli.db'}"]
    result = runner.invoke(cli, [*db, "compare", str(real), str(same), "--min-score", "99"])
    assert result.exit_code == 0, result.output
    assert "1:1" in result.output
    result = runner.invoke(cli, [*db, "compare", str(real), str(other), "--min-score", "99"])
    assert result.exit_code == 1
    result = runner.invoke(cli, [*db, "compare", str(real), str(other), "--json"])
    assert json.loads(result.output)["verdict"] == "different"
    assert runner.invoke(cli, [*db, "convert-engines"]).exit_code == 0


@requires_libreoffice
def test_cli_convert(tmp_path):
    source = tmp_path / "act.docx"
    source.write_bytes(word_document())
    db = ["--db", f"sqlite:///{tmp_path / 'cli.db'}"]
    result = runner.invoke(cli, [*db, "convert", str(source)])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "act.pdf").is_file()
    assert "Against the original" in result.output
    again = runner.invoke(cli, [*db, "convert", str(source), "--no-score"])
    assert again.exit_code == 0 and (tmp_path / "act-2.pdf").is_file()  # never over a file
