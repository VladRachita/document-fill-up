from __future__ import annotations

import io

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from docfill.config import Settings
from docfill.readers.ocr import tesseract_available

FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def _spacy_model_installed(name: str = "en_core_web_sm") -> bool:
    try:
        import spacy

        return spacy.util.is_package(name)
    except ImportError:
        return False


requires_tesseract = pytest.mark.skipif(
    not tesseract_available(), reason="Tesseract OCR is not installed"
)
requires_spacy_model = pytest.mark.skipif(
    not _spacy_model_installed(), reason="spaCy model en_core_web_sm is not installed"
)

ID_CARD_LINES = [
    "IDENTITY CARD",
    "Surname: SMITH",
    "Given names: John Michael",
    "Place of birth: Manchester",
    "Address: 12 Baker Street, London NW1 6XE, United Kingdom",
]


@pytest.fixture
def settings(tmp_path) -> Settings:
    """Settings with a throw-away database and NER disabled (fast, deterministic)."""
    return Settings(database_url=f"sqlite:///{tmp_path / 'test.db'}", spacy_model="")


@pytest.fixture
def ner_settings(tmp_path) -> Settings:
    return Settings(database_url=f"sqlite:///{tmp_path / 'test.db'}")


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=0, cols=len(table[0]))
        for row in table:
            cells = grid.add_row().cells
            for cell, text in zip(cells, row, strict=True):
                cell.text = text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_text_pdf(pages: list[list[str]]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            pdf.drawString(60, y, line)
            y -= 20
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_image(lines: list[str], fmt: str = "PNG") -> bytes:
    try:
        font = ImageFont.truetype(FONT_PATH, 28)
    except OSError:
        font = ImageFont.load_default(size=28)
    image = Image.new("RGB", (1200, 90 + 70 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        draw.text((40, 40 + 70 * index), line, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format=fmt)
    return buffer.getvalue()


def make_scanned_pdf(lines: list[str]) -> bytes:
    """A PDF containing only a picture of text (no text layer), like a scanner produces."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    image = ImageReader(io.BytesIO(make_image(lines, "JPEG")))
    pdf.drawImage(image, 20, 450, width=560, height=560 * 440 / 1200)
    pdf.save()
    return buffer.getvalue()


def make_pdf_form(fields: list[str]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 780
    for name in fields:
        pdf.drawString(60, y + 5, f"{name}:")
        pdf.acroForm.textfield(name=name, x=200, y=y, width=250, height=20, borderWidth=1)
        y -= 40
    pdf.save()
    return buffer.getvalue()


@pytest.fixture
def id_card_png() -> bytes:
    return make_image(ID_CARD_LINES)
