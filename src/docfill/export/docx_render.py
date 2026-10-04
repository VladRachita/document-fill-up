"""Render a filled text standard document as a Word document (python-docx), for the documents
the filer edits before signing (``output: docx``, e.g. the act constitutiv).

The layout of the model filers write the act in (model-act-constitutiv-srl-asociat-unic.docx,
next to the standard documents): Times New Roman 12 on a Letter page with 2.54 cm margins,
justified paragraphs; ``# Title`` lines centred in bold, the first one larger; ``## Chapter``
lines centred in bold; the number of an article (``Art. 1.2.``) in bold; list items (``— ...``)
close together; the date and the signature at the end, left aligned, ``(semnătura)`` in
italics. The lines follow the same small language as the PDF renderer
(:mod:`docfill.templates.markup`): blank lines separate paragraphs, single line breaks are kept.
The values of the fields a standard document names in ``bold`` are written in bold (the firm,
the associate and the administrator in the act, as in the model).
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable, Mapping
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

FONT = "Times New Roman"
SIZE = 12
TITLE_SIZE = 14
LINE_SPACING = 1.15
_ARTICLE = re.compile(r"^(Art\.\s*\d+(?:\.\d+)*\.)")
_LIST_ITEM = re.compile(r"^\s*[—–-]\s")
_DATE = re.compile(r"^\s*Data\s*:")
_RULE = re.compile(r"^\s*_{5,}\s*$")
_NOTE = re.compile(r"^\s*\(.*\)\s*$")


def _base_document(metadata: Mapping[str, str]):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(SIZE)
    normal.paragraph_format.space_after = Pt(0)
    defaults = document.styles.element.find(qn("w:docDefaults"))
    for properties in (normal.element.get_or_add_rPr(), *defaults.iter(qn("w:rPr"))):
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = properties.get_or_add_rFonts()
        for attribute in list(fonts.attrib):  # a theme font would win over the named one
            if attribute.lower().endswith("theme"):
                del fonts.attrib[attribute]
        for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attribute), FONT)
        language = properties.find(qn("w:lang"))
        if language is None:
            language = OxmlElement("w:lang")
            properties.append(language)
        language.set(qn("w:val"), "ro-RO")
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    properties = document.core_properties
    properties.title = metadata.get("title", "")
    properties.subject = metadata.get("subject", "")
    properties.keywords = metadata.get("keywords", "")
    properties.author = metadata.get("author", "docfill")
    properties.last_modified_by = "docfill"
    properties.comments = ""
    properties.created = properties.modified = datetime.now(timezone.utc).replace(microsecond=0)
    return document


def _paragraph(
    document,
    alignment: WD_ALIGN_PARAGRAPH,
    before: float = 0,
    after: float = 6,
    line: float | None = LINE_SPACING,
) -> Paragraph:
    paragraph = document.add_paragraph()
    layout = paragraph.paragraph_format
    layout.alignment = alignment
    layout.space_before, layout.space_after = Pt(before), Pt(after)
    if line:
        layout.line_spacing = line
    return paragraph


def _write(
    paragraph: Paragraph, text: str, bold: bool = False, italic: bool = False, size: int = SIZE
) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def _rule(paragraph: Paragraph) -> None:
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _write_line(paragraph: Paragraph, segments: list[dict], bold: set[str], note: bool) -> None:
    """One line of the document: the number of an article and the values of ``bold`` fields in
    bold, a note (``(semnătura)``) in italics."""
    text = "".join(segment["text"] for segment in segments)
    article = _ARTICLE.match(text)
    position = 0
    for segment in segments:
        piece, strong = segment["text"], segment.get("field") in bold
        if article and position < article.end() and not segment.get("field"):
            cut = article.end() - position
            if piece[:cut]:
                _write(paragraph, piece[:cut], bold=True)
            position += len(piece)
            piece = piece[cut:]
        else:
            position += len(piece)
        if piece:
            _write(paragraph, piece, bold=strong, italic=note)


def render_text_docx(
    lines: Iterable[Mapping[str, Any]],
    metadata: Mapping[str, str] | None = None,
    bold: Iterable[str] = (),
) -> bytes:
    """``lines``: the filled document as :func:`docfill.templates.placeholders.preview_lines`
    gives it (the kind of every line and its text, with the values marked by their field);
    ``bold``: the fields whose values are written in bold."""
    document = _base_document(dict(metadata or {}))
    strong = set(bold)
    paragraph_lines: list[list[dict]] = []  # the lines of the paragraph being written
    titles = 0
    closing = 0  # paragraphs written since "Data:": the signature block

    def flush() -> None:
        nonlocal closing
        if not paragraph_lines:
            return
        texts = ["".join(segment["text"] for segment in line) for line in paragraph_lines]
        first = texts[0]
        if _DATE.match(first):
            closing = 1
            paragraph = _paragraph(document, WD_ALIGN_PARAGRAPH.LEFT, 12, 3, None)
        elif closing:
            closing += 1
            after = 3 if _RULE.match(first) or _NOTE.match(first) else 12
            paragraph = _paragraph(document, WD_ALIGN_PARAGRAPH.LEFT, 0, after, None)
        elif _LIST_ITEM.match(first) and not first.rstrip().endswith(":"):
            paragraph = _paragraph(document, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2)
        else:
            paragraph = _paragraph(document, WD_ALIGN_PARAGRAPH.JUSTIFY)
        for index, (line, text) in enumerate(zip(paragraph_lines, texts, strict=True)):
            if index and paragraph.runs:
                paragraph.runs[-1].add_break()
            _write_line(paragraph, line, strong, closing > 0 and bool(_NOTE.match(text)))
        paragraph_lines.clear()

    for line in lines:
        kind, segments = line["kind"], list(line.get("segments") or [])
        if kind == "text":
            paragraph_lines.append(segments)
            continue
        flush()
        text = "".join(segment["text"] for segment in segments)
        if kind == "h1":
            titles += 1
            first = titles == 1
            paragraph = _paragraph(document, WD_ALIGN_PARAGRAPH.CENTER, 0, 3 if first else 10, None)
            _write(paragraph, text, bold=True, size=TITLE_SIZE if first else SIZE)
        elif kind == "h2":
            paragraph = _paragraph(document, WD_ALIGN_PARAGRAPH.CENTER, 10, 6)
            _write(paragraph, text, bold=True)
        elif kind == "hr":
            _rule(_paragraph(document, WD_ALIGN_PARAGRAPH.LEFT, 0, 6, None))
    flush()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
