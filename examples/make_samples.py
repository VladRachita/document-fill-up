"""Generate sample input documents to try docfill with.

python examples/make_samples.py
docfill extract examples/samples/*

examples/samples/learning/ holds a set to watch docfill learn: fictitious people, each as a
clean identity card and as bad scans (tilted photo, low resolution, faded, glare), with the
correct values in ANSWERS.md. See docs/LOCAL_TESTING.md.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import date
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from docfill.samples import Person, ro_id_card, ro_id_card_jpeg

OUT = Path(__file__).parent / "samples"


def _font(size: int) -> ImageFont.ImageFont:
    for path in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "arial.ttf"):
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def id_card_image() -> Image.Image:
    lines = [
        "IDENTITY CARD",
        "Surname: SMITH",
        "Given names: John Michael",
        "Place of birth: Manchester",
        "Address: 12 Baker Street, London NW1 6XE, United Kingdom",
    ]
    image = Image.new("RGB", (1200, 90 + 70 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        draw.text((40, 40 + 70 * index), line, fill="black", font=_font(28))
    return image


# --------------------------------------------------------------------------- learning set

# Fictitious people (valid CNP and MRZ). Each has something worth checking: diacritics lost by
# OCR, a small commune docfill does not know, an expired card, a holder under 18.
PEOPLE = {
    "popescu": Person(),
    "stefanescu": Person(
        last_name="ȘTEFĂNESCU",
        first_name="ANA-MARIA",
        sex="F",
        birth=date(1994, 3, 8),
        birth_county_code="SB",
        birth_locality="Mun.Mediaș",
        county_code="TM",
        locality="Com.Săcălaz",
        street_line="Str.Principală nr.12",
        series="TZ",
        number="604213",
        issued_by="SPCLEP Timișoara",
        issued=date(2021, 3, 1),
        expires=date(2031, 3, 8),
        serial="457",
    ),
    "muresan": Person(
        last_name="MUREȘAN",
        first_name="VLAD",
        birth=date(1979, 7, 21),
        birth_county_code="AB",
        birth_locality="Mun.Aiud",
        county_code="AB",
        locality="Mun.Alba Iulia",
        street_line="Str.Mihai Viteazul nr.12 bl.A2 sc.1 et.3 ap.10",
        series="AX",
        number="781245",
        issued_by="SPCLEP Alba Iulia",
        issued=date(2015, 7, 21),
        expires=date(2025, 7, 21),  # expired: the wizard must say so
        serial="088",
    ),
    "dumitru": Person(
        last_name="DUMITRU",
        first_name="ELENA",
        sex="F",
        birth=date(2009, 5, 2),  # under 18: the PFA / II age check must fail
        birth_county_code="CJ",
        birth_locality="Mun.Cluj-Napoca",
        county_code="CJ",
        locality="Mun.Turda",
        street_line="Str.Republicii nr.3",
        series="KX",
        number="112233",
        issued_by="SPCLEP Turda",
        issued=date(2023, 5, 10),
        expires=date(2027, 5, 2),
        serial="021",
    ),
}


def _tilted_photo(image: Image.Image) -> Image.Image:
    rotated = image.rotate(4, expand=True, fillcolor=(90, 90, 90), resample=Image.BICUBIC)
    return rotated.filter(ImageFilter.GaussianBlur(1.6))


def _low_resolution(image: Image.Image) -> Image.Image:
    return image.resize((image.width * 38 // 100, image.height * 38 // 100), Image.BILINEAR)


def _faded(image: Image.Image) -> Image.Image:
    return ImageEnhance.Brightness(ImageEnhance.Contrast(image).enhance(0.45)).enhance(1.25)


def _glare(image: Image.Image) -> Image.Image:
    mask = Image.new("L", image.size, 0)
    ImageDraw.Draw(mask).ellipse((650, 300, 1250, 700), fill=217)
    white = Image.new("RGB", image.size, (255, 255, 255))
    return Image.composite(white, image, mask.filter(ImageFilter.GaussianBlur(60)))


# variant -> (how the card is damaged, JPEG quality)
SCANS: dict[str, tuple[Callable[[Image.Image], Image.Image], int]] = {
    "clean": (lambda image: image, 88),
    "photo": (_tilted_photo, 88),
    "lowres": (_low_resolution, 30),
    "faded": (_faded, 88),
    "glare": (_glare, 88),
}


def _answers(person: Person) -> list[tuple[str, str]]:
    return [
        ("last_name", person.last_name),
        ("first_name", person.first_name),
        ("cnp", person.cnp),
        ("sex", person.sex),
        ("date_of_birth", f"{person.birth:%d.%m.%Y}"),
        ("place_of_birth", person.birth_locality.replace(".", ". ", 1)),
        ("id_series / id_number", f"{person.series} / {person.number}"),
        ("id_issued_by", person.issued_by),
        ("id_issue_date / id_expiry_date", f"{person.issued:%d.%m.%Y} / {person.expires:%d.%m.%Y}"),
        ("city", person.locality.replace(".", ". ", 1)),
        (
            "domicile (as printed)",
            f"Jud.{person.county_code} {person.locality}, {person.street_line}",
        ),
    ]


def client_sheet(person: Person) -> Document:
    """A client sheet with a label docfill does not know ("Localitatea natală"): type the place
    of birth once in the wizard and the label is learned for the next sheet."""
    document = Document()
    document.add_heading("Fișă client", 1)
    document.add_paragraph(f"Nume: {person.last_name}")
    document.add_paragraph(f"Prenume: {person.first_name}")
    document.add_paragraph(f"Localitatea natală: {person.birth_locality.split('.', 1)[-1]}")
    return document


def learning_set(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Correct values of the learning set",
        "",
        "Every `ci_*.jpg` is a fictitious person's identity card; the suffix says how it was",
        "damaged: `clean`, `photo` (tilted, out of focus), `lowres` (small, heavy JPEG), `faded`",
        "(low contrast), `glare` (a light reflection over the names). `fisa_*.docx` are client",
        "sheets with a label docfill does not know. Compare what docfill proposes with these",
        "values (upper / lower case does not matter), correct it in the wizard and watch",
        "`docfill learn stats`. See docs/LOCAL_TESTING.md.",
    ]
    for key, person in PEOPLE.items():
        card = ro_id_card(person)
        for variant, (damage, quality) in SCANS.items():
            buffer = io.BytesIO()
            damage(card).convert("RGB").save(buffer, format="JPEG", quality=quality)
            (out / f"ci_{key}_{variant}.jpg").write_bytes(buffer.getvalue())
        client_sheet(person).save(out / f"fisa_{key}.docx")
        lines += ["", f"## {key}", "", "| Field | Correct value |", "|---|---|"]
        lines += [f"| {name} | {value} |" for name, value in _answers(person)]
    (out / "ANSWERS.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(exist_ok=True)

    # 0. A Romanian identity card of a fictitious person (CNP, MRZ, domicile...)
    (OUT / "ci_popescu.jpg").write_bytes(ro_id_card_jpeg(Person()))

    # 1. A photographed / scanned ID card (JPEG) -> OCR
    id_card_image().save(OUT / "id_card.jpg", quality=90)

    # 2. The same card inside a PDF without a text layer (scanner output) -> OCR
    buffer = io.BytesIO()
    id_card_image().save(buffer, format="PNG")
    pdf = canvas.Canvas(str(OUT / "scanned_id_card.pdf"), pagesize=A4)
    pdf.drawImage(ImageReader(io.BytesIO(buffer.getvalue())), 20, 450, width=560, height=205)
    pdf.save()

    # 3. A Romanian application form in Word with a key/value table
    document = Document()
    document.add_heading("Cerere / Application", 1)
    document.add_paragraph("Vă rog să completați datele de mai jos.")
    table = document.add_table(rows=0, cols=2)
    for key, value in [
        ("Nume", "POPESCU"),
        ("Prenume", "Ion Andrei"),
        ("Domiciliu", "Str. Florilor nr. 5, 400001 Cluj-Napoca, jud. Cluj, România"),
        ("Locul nașterii", "Făgăraș"),
        ("IBAN", "RO49 AAAA 1B31 0075 9384 0000"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    document.save(OUT / "application_ro.docx")

    # 4. A free-text letter (no labels) -> ML named-entity recognition
    pdf = canvas.Canvas(str(OUT / "letter.pdf"), pagesize=A4)
    text = pdf.beginText(60, 780)
    for line in [
        "Dear Sir or Madam,",
        "",
        "My name is Maria Garcia and I was born in Madrid. I currently live in Berlin,",
        "Germany, at Hauptstraße 5. Please update your records accordingly.",
        "",
        "Kind regards,",
        "Maria Garcia",
    ]:
        text.textLine(line)
    pdf.drawText(text)
    pdf.save()

    # 5. The learning set: several people, each clean and as bad scans, with the answers
    learning_set(OUT / "learning")

    for path in sorted(OUT.rglob("*")):
        if path.is_file():
            print(path)


if __name__ == "__main__":
    main()
